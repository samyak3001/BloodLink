import os
import re
import html
import subprocess
import pypdf
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import generate_complete_55page_report as orig

def build_custom_pages():
    raw_pages = orig.get_html_pages()
    pages = list(raw_pages) # Make a mutable copy

    # -------------------------------------------------------------------------
    # PAGE 1: TITLE PAGE (Exact Match to Template)
    # -------------------------------------------------------------------------
    pages[0] = """
    <div style="text-align: center; padding-top: 25px;">
      <h1 style="font-size: 16pt; font-weight: bold; margin-bottom: 12px; line-height: 1.3; text-transform: uppercase;">
        BLOODLINK: REAL-TIME EMERGENCY BLOOD DONOR &amp; HOSPITAL MATCHING PLATFORM
      </h1>
      <p style="font-size: 13pt; font-weight: bold; margin-bottom: 25px;">
        A PROJECT REPORT
      </p>
      <p style="font-size: 12pt; font-style: italic; margin-bottom: 25px;">
        Submitted by
      </p>
      <p style="font-size: 14pt; font-weight: bold; margin-bottom: 4px;">
        PON SAMYAK M S
      </p>
      <p style="font-size: 11pt; font-weight: bold; margin-bottom: 30px;">
        (Reg. No.: 310624148061 / 310622104088)
      </p>
      <p style="font-size: 12pt; font-style: italic; margin-bottom: 12px;">
        in partial fulfillment for the award of the Degree<br>of
      </p>
      <p style="font-size: 13pt; font-weight: bold; margin-bottom: 12px;">
        BACHELOR OF ENGINEERING
      </p>
      <p style="font-size: 12pt; font-style: italic; margin-bottom: 12px;">
        in
      </p>
      <p style="font-size: 13pt; font-weight: bold; margin-bottom: 25px;">
        COMPUTER SCIENCE AND ENGINEERING (AI &amp; ML)
      </p>

      <div style="margin: 15px auto; width: 100px; height: 100px; border: 3px double #1e3a8a; border-radius: 50%; display: flex; flex-direction: column; align-items: center; justify-content: center; background: #f8fafc;">
        <div style="font-size: 7pt; font-weight: bold; color: #1e3a8a; text-align: center; line-height: 1.1;">
          EASWARI<br>ENGG COLLEGE<br>
          <span style="font-size: 6pt; color: #dc2626;">AUTONOMOUS</span><br>
          <span style="font-size: 5.5pt; color: #15803d;">SRM GROUP</span>
        </div>
      </div>

      <div style="margin-top: 25px;">
        <p style="font-size: 13pt; font-weight: bold; margin-bottom: 4px;">
          EASWARI ENGINEERING COLLEGE
        </p>
        <p style="font-size: 10pt; font-weight: bold; margin-bottom: 2px;">
          (AN AUTONOMOUS INSTITUTION) – AFFILIATED TO ANNA UNIVERSITY,
        </p>
        <p style="font-size: 11pt; font-weight: bold; margin-bottom: 20px;">
          CHENNAI
        </p>
        <p style="font-size: 12pt; font-weight: bold;">
          OCTOBER 2026
        </p>
      </div>
    </div>
    """

    # -------------------------------------------------------------------------
    # PAGE 2: BONAFIDE CERTIFICATE (Exact Match to Template)
    # -------------------------------------------------------------------------
    pages[1] = """
    <div style="text-align: center; padding-top: 40px; margin-bottom: 40px;">
      <h2 style="font-size: 14pt; font-weight: bold; margin-bottom: 50px;">EASWARI ENGINEERING COLLEGE</h2>
      <h1 style="font-size: 14pt; font-weight: bold; text-decoration: none; margin-bottom: 40px;">BONAFIDE CERTIFICATE</h1>
    </div>
    <p style="text-align: justify; line-height: 2.0; font-size: 12pt; margin-bottom: 80px;">
      Certified that this project report titled <strong>“BLOODLINK: REAL-TIME EMERGENCY BLOOD DONOR &amp; HOSPITAL MATCHING PLATFORM”</strong> is the Bonafide work of <strong>“PON SAMYAK M S (310624148061)”</strong>, who carried out the project work under my supervision.
    </p>
    <div style="margin-top: 220px; display: flex; justify-content: space-between; font-size: 11pt; font-weight: bold;">
      <div style="width: 45%; text-align: left;">
        HEAD OF THE DEPARTMENT
      </div>
      <div style="width: 45%; text-align: right;">
        SUPERVISOR
      </div>
    </div>
    """

    # -------------------------------------------------------------------------
    # PAGE 3: CERTIFICATE OF VIVA-VOCE EXAMINATION (Exact Match to Template)
    # -------------------------------------------------------------------------
    pages[2] = """
    <div style="text-align: center; padding-top: 40px; margin-bottom: 30px;">
      <h2 style="font-size: 14pt; font-weight: bold; margin-bottom: 40px;">EASWARI ENGINEERING COLLEGE</h2>
      <h1 style="font-size: 13pt; font-weight: bold; margin-bottom: 40px;">Certificate of Viva-voce-Examination</h1>
    </div>
    <p style="text-align: justify; line-height: 1.8; font-size: 11.5pt; margin-bottom: 50px;">
      This is to certify that the Viva Voce Examination of the Project Work entitled <strong>“BLOODLINK: REAL-TIME EMERGENCY BLOOD DONOR &amp; HOSPITAL MATCHING PLATFORM”</strong> submitted by <strong>Mr. PON SAMYAK M S (Reg. No. 310624148061)</strong> in partial fulfillment of the requirements for the award of the degree of <strong>BACHELOR OF ENGINEERING in COMPUTER SCIENCE AND ENGINEERING (AI &amp; ML)</strong> was conducted on ________________ (Date) in the Department of Computer Science and Engineering (AI &amp; ML), Easwari Engineering College, Ramapuram – Chennai.
    </p>
    <div style="margin-top: 90px; display: flex; justify-content: space-between; font-size: 11pt; line-height: 1.8;">
      <div style="width: 48%; text-align: left;">
        <p style="font-weight: bold; margin-bottom: 18px;">Internal Examiner</p>
        <p>Name :</p>
        <p>Designation :</p>
        <p>Address :</p>
        <p style="margin-top: 35px;">Date :</p>
      </div>
      <div style="width: 48%; text-align: left;">
        <p style="font-weight: bold; margin-bottom: 18px;">External Examiner</p>
        <p>Name :</p>
        <p>Designation :</p>
        <p>Address :</p>
      </div>
    </div>
    """

    # -------------------------------------------------------------------------
    # PAGE 5: ACKNOWLEDGEMENT (Exact Match to Template)
    # -------------------------------------------------------------------------
    pages[4] = """
    <div style="text-align: center; padding-top: 30px; margin-bottom: 35px;">
      <h1 style="font-size: 14pt; font-weight: bold; text-transform: uppercase;">ACKNOWLEDGEMENT</h1>
    </div>
    <p style="text-align: justify; line-height: 1.8; font-size: 11.5pt; margin-bottom: 22px;">
      We express our sincere gratitude to the Principal <strong>Dr. P. Deiva Sundari</strong> and the Management of our college for providing us with the opportunity to carry out this project work.
    </p>
    <p style="text-align: justify; line-height: 1.8; font-size: 11.5pt; margin-bottom: 22px;">
      We convey our heartfelt thanks to the Head of the Department <strong>Dr. Y. Justin Dhas</strong>, Department of Computer Science and Engineering (AI &amp; ML), for the continuous encouragement and support.
    </p>
    <p style="text-align: justify; line-height: 1.8; font-size: 11.5pt; margin-bottom: 22px;">
      We are extremely thankful to our Project Guide <strong>Ms. Divyanjali</strong>, Department of Computer Science and Engineering (AI &amp; ML) for valuable guidance, suggestions, and motivation throughout the completion of this project.
    </p>
    <p style="text-align: justify; line-height: 1.8; font-size: 11.5pt; margin-bottom: 22px;">
      We would also like to extend our gratitude to all the faculty members of the Department of Computer Science and Engineering (AI &amp; ML) for their support.
    </p>
    <p style="text-align: justify; line-height: 1.8; font-size: 11.5pt; margin-bottom: 22px;">
      Finally, we thank our parents and friends for their constant encouragement and help during the course of this project.
    </p>
    """

    # -------------------------------------------------------------------------
    # PAGE 6: TABLE OF CONTENTS (Exact Match to Template Page Numbers)
    # -------------------------------------------------------------------------
    pages[5] = """
    <div style="text-align: center; padding-top: 15px; margin-bottom: 15px;">
      <h1 style="font-size: 14pt; font-weight: bold; text-transform: uppercase;">TABLE OF CONTENTS</h1>
    </div>
    <table style="width: 100%; border-collapse: collapse; font-size: 9.5pt; line-height: 1.28;">
      <thead>
        <tr style="border-bottom: 2px solid #000;">
          <th style="width: 15%; text-align: left; padding: 4px 0;">CHAPTER NO.</th>
          <th style="width: 73%; text-align: left; padding: 4px 0;">TITLE</th>
          <th style="width: 12%; text-align: right; padding: 4px 0;">PAGE No.</th>
        </tr>
      </thead>
      <tbody>
        <tr><td></td><td>ABSTRACT</td><td style="text-align: right;">4</td></tr>
        <tr><td></td><td>ACKNOWLEDGEMENT</td><td style="text-align: right;">5</td></tr>
        <tr style="font-weight: bold;"><td>1</td><td>INTRODUCTION</td><td style="text-align: right;">7</td></tr>
        <tr><td>1.1</td><td>OVERVIEW</td><td style="text-align: right;">7</td></tr>
        <tr><td>1.2</td><td>PROBLEM DEFINITION</td><td style="text-align: right;">8</td></tr>
        <tr><td>1.3</td><td>OBJECTIVES OF THE SYSTEM</td><td style="text-align: right;">8</td></tr>
        <tr><td>1.4</td><td>SCOPE OF THE PROJECT</td><td style="text-align: right;">9</td></tr>
        <tr><td>1.5</td><td>MOTIVATION</td><td style="text-align: right;">10</td></tr>
        <tr><td>1.6</td><td>HEMATOLOGICAL COMPATIBILITY DYNAMICS &amp; INFORMATICS</td><td style="text-align: right;">11</td></tr>
        <tr><td>1.7</td><td>EMERGENCY DISPATCH ACCESSIBILITY &amp; REGIONAL NOTIFICATIONS</td><td style="text-align: right;">12</td></tr>
        <tr style="font-weight: bold;"><td>2</td><td>RELATED WORK &amp; LITERATURE REVIEW</td><td style="text-align: right;">13</td></tr>
        <tr><td>2.1</td><td>LITERATURE REVIEW</td><td style="text-align: right;">13</td></tr>
        <tr><td>2.2</td><td>EXISTING SYSTEM</td><td style="text-align: right;">16</td></tr>
        <tr><td>2.3</td><td>PROPOSED SYSTEM</td><td style="text-align: right;">17</td></tr>
        <tr style="font-weight: bold;"><td>3</td><td>SYSTEM ANALYSIS &amp; FEASIBILITY STUDY</td><td style="text-align: right;">19</td></tr>
        <tr><td>3.1</td><td>SYSTEM REQUIREMENTS</td><td style="text-align: right;">19</td></tr>
        <tr><td>3.2</td><td>SOFTWARE AND HARDWARE REQUIREMENTS</td><td style="text-align: right;">20</td></tr>
        <tr><td>3.3</td><td>COMPREHENSIVE FEASIBILITY ANALYSIS</td><td style="text-align: right;">21</td></tr>
        <tr><td>3.4</td><td>USE CASE ANALYSIS &amp; OPERATIONAL ACTOR-GOAL MATRIX</td><td style="text-align: right;">23</td></tr>
        <tr style="font-weight: bold;"><td>4</td><td>SYSTEM DESIGN &amp; ARCHITECTURE</td><td style="text-align: right;">24</td></tr>
        <tr><td>4.1</td><td>SYSTEM ARCHITECTURE</td><td style="text-align: right;">24</td></tr>
        <tr><td>4.2</td><td>MODULE DESCRIPTION</td><td style="text-align: right;">25</td></tr>
        <tr><td>4.3</td><td>DATABASE SCHEMA AND COMPOUND INDEXING DESIGN</td><td style="text-align: right;">26</td></tr>
        <tr><td>4.4</td><td>CLASS DIAGRAM &amp; DOMAIN MODELLING</td><td style="text-align: right;">27</td></tr>
        <tr><td>4.5</td><td>SEQUENCE DIAGRAMS</td><td style="text-align: right;">28</td></tr>
        <tr><td>4.6</td><td>DATA FLOW DIAGRAMS (DFD)</td><td style="text-align: right;">30</td></tr>
        <tr><td>4.7</td><td>ENTITY-RELATIONSHIP ARCHITECTURE AND DATA DICTIONARY</td><td style="text-align: right;">31</td></tr>
        <tr style="font-weight: bold;"><td>5</td><td>IMPLEMENTATION &amp; CODE ANALYSIS</td><td style="text-align: right;">32</td></tr>
        <tr><td>5.1</td><td>IMPLEMENTATION OVERVIEW</td><td style="text-align: right;">32</td></tr>
        <tr><td>5.2</td><td>KEY FUNCTIONS (CODE EXPLANATION)</td><td style="text-align: right;">32</td></tr>
        <tr><td>5.3</td><td>SYSTEM VERIFICATION, TESTING AND QUALITY ASSURANCE</td><td style="text-align: right;">35</td></tr>
        <tr style="font-weight: bold;"><td>6</td><td>RESULTS AND DISCUSSION</td><td style="text-align: right;">37</td></tr>
        <tr><td>6.1</td><td>PERFORMANCE AND LATENCY BENCHMARKS</td><td style="text-align: right;">37</td></tr>
        <tr><td>6.2</td><td>COMPARATIVE EVALUATION WITH EXISTING SYSTEMS</td><td style="text-align: right;">38</td></tr>
        <tr><td>6.3</td><td>EMPIRICAL USABILITY EVALUATION (SYSTEM USABILITY SCALE)</td><td style="text-align: right;">38</td></tr>
        <tr style="font-weight: bold;"><td>7</td><td>CONCLUSION &amp; FUTURE WORK</td><td style="text-align: right;">39</td></tr>
        <tr><td>7.1</td><td>CONCLUSION</td><td style="text-align: right;">39</td></tr>
        <tr><td>7.2</td><td>FUTURE ENHANCEMENTS</td><td style="text-align: right;">39</td></tr>
        <tr style="font-weight: bold;"><td></td><td>APPENDIX (SOURCE CODE)</td><td style="text-align: right;">41</td></tr>
        <tr style="font-weight: bold;"><td></td><td>SCREENSHOTS OF THE PROJECT</td><td style="text-align: right;">51</td></tr>
        <tr style="font-weight: bold;"><td></td><td>REFERENCES</td><td style="text-align: right;">55</td></tr>
      </tbody>
    </table>
    """

    # -------------------------------------------------------------------------
    # Re-structure Appendix so it occupies Pages 41 to 50 (10 pages)
    # Page 41: Project Structure & Config
    # Page 42: Database Connection
    # Page 43: User & Donor Models
    # Page 44: Hospital & Request Models
    # Page 45: Matching Service
    # Page 46: Emergency Request Controller
    # Page 47: Socket Gateway
    # Page 48: Frontend API Client & View
    # Page 49: REST API Specification Examples
    # Page 50: Automated Integration Test Suite
    # -------------------------------------------------------------------------
    
    # Page 41: Project Structure, Monorepo Configuration & Environment Variables
    pages[40] = """
    <div style="text-align: center; padding-top: 10px; margin-bottom: 12px;">
      <h1 style="font-size: 13pt; font-weight: bold; text-transform: uppercase;">APPENDIX<br>SOURCE CODE &amp; ARCHITECTURAL SPECIFICATIONS</h1>
    </div>
    <h3 style="font-size: 10.5pt; font-weight: bold; margin-bottom: 6px;">1. Project Monorepo Structure &amp; Directory Tree</h3>
    <div style="background: #f8fafc; border: 1px solid #cbd5e1; padding: 8px; margin-bottom: 10px; font-family: 'Courier New', monospace; font-size: 7.5pt; line-height: 1.2;">
BloodLink/
├── client/                     # Frontend React 18 + Vite SPA
│   ├── src/
│   │   ├── api/                # Axios API clients (auth, donor, hospital, requests)
│   │   ├── components/         # Shared UI primitives & EmergencyRequestCard
│   │   ├── context/            # AuthContext, NotificationContext, SocketContext
│   │   ├── hooks/              # useAuth, useSocket, useNotifications
│   │   ├── pages/              # DonorDashboard, HospitalDashboard, AdminAudit
│   │   └── routes/             # AppRoutes (RBAC Route Guards)
│   ├── tailwind.config.js      # Healthcare semantic design tokens
│   └── vite.config.ts          # Vite build config with backend proxy
├── server/                     # Backend Node.js + Express API & WebSocket Gateway
│   ├── src/
│   │   ├── config/             # MongoDB connection, blood compatibility rules
│   │   ├── controllers/        # Auth, Donor, Hospital, EmergencyRequest, Admin
│   │   ├── middleware/         # authenticate, authorize, errorHandler, rateLimiter
│   │   ├── models/             # User, DonorProfile, HospitalProfile, EmergencyRequest
│   │   ├── routes/             # REST routes mounted under /api
│   │   ├── services/           # matchingService, socketService
│   │   └── tests/              # Vitest automated regression test suites
│   └── package.json
└── docker-compose.yml          # Containerized MongoDB 7.0 deployment
    </div>
    <h3 style="font-size: 10.5pt; font-weight: bold; margin-bottom: 6px;">2. Environment Variables Specification (server/.env)</h3>
    <div style="background: #f8fafc; border: 1px solid #cbd5e1; padding: 8px; font-family: 'Courier New', monospace; font-size: 7.5pt; line-height: 1.2;">
PORT=5000
NODE_ENV=production
MONGODB_URI=mongodb+srv://&lt;username&gt;:&lt;password&gt;@cluster0.mongodb.net/bloodlink?retryWrites=true&amp;w=majority
JWT_SECRET=&lt;your_secure_jwt_secret_key&gt;
JWT_EXPIRE=24h
CLIENT_URL=http://localhost:5173
    </div>
    """

    # Page 48: Frontend API Client & Donor Requests View Component
    pages[47] = """
    <h2 class="section-title">SOURCE CODE 8: FRONTEND API CLIENT &amp; REAL-TIME VIEW</h2>
    <p class="justify">
      <strong>File: client/src/api/requestsApi.ts &amp; DonorRequestsPage.tsx</strong>
    </p>
    <div style="background: #f8fafc; border: 1px solid #cbd5e1; padding: 8px; margin: 8px 0; font-family: 'Courier New', monospace; font-size: 7pt; line-height: 1.18;">
// Fetch matching emergencies for donor
export async function fetchDonorEmergencyRequests(token: string): Promise&lt;EmergencyRequestItem[]&gt; {
  const res = await fetch(`${getApiBaseUrl()}/api/donors/requests`, {
    headers: { 'Authorization': `Bearer ${token}`, 'Content-Type': 'application/json' }
  });
  if (!res.ok) throw new Error('Failed to fetch emergency requests.');
  const json = await res.json();
  return Array.isArray(json) ? json : (json.data || []);
}

// Single-click emergency acceptance API call
export async function acceptEmergencyRequestApi(requestId: string, token: string): Promise&lt;any&gt; {
  const res = await fetch(`${getApiBaseUrl()}/api/requests/${requestId}/accept`, {
    method: 'POST',
    headers: { 'Authorization': `Bearer ${token}`, 'Content-Type': 'application/json' }
  });
  if (!res.ok) {
    const err = await res.json().catch(() =&gt; ({}));
    throw new Error(err.message || 'Acceptance failed. Request may already be fulfilled.');
  }
  return res.json();
}

// React Functional Component: Donor Emergency Dispatch View
export const DonorRequestsView: React.FC = () =&gt; {
  const { token } = useAuth();
  const [requests, setRequests] = useState&lt;EmergencyRequestItem[]&gt;([]);
  const [loading, setLoading] = useState(true);

  useEffect(() =&gt; {
    if (token) fetchDonorEmergencyRequests(token).then(setRequests).finally(() =&gt; setLoading(false));
  }, [token]);

  if (loading) return &lt;LoadingSpinner message="Scanning regional emergencies..." /&gt;;
  return (
    &lt;div className="space-y-4"&gt;
      &lt;h2 className="text-xl font-bold"&gt;Active Matching Emergencies ({requests.length})&lt;/h2&gt;
      {requests.map(req =&gt; &lt;EmergencyCard key={req._id} request={req} /&gt;)}
    &lt;/div&gt;
  );
};
    </div>
    """

    # Page 49: REST API Specification Examples
    pages[48] = """
    <h2 class="section-title">REST API SPECIFICATION REQUEST &amp; RESPONSE ENVELOPES</h2>
    <div style="background: #f8fafc; border: 1px solid #cbd5e1; padding: 8px; margin: 8px 0; font-family: 'Courier New', monospace; font-size: 7pt; line-height: 1.18;">
// 1. REQUEST: POST /api/auth/register (Donor Self-Service Onboarding)
{
  "name": "Samyak Pon",
  "email": "donor@example.com",
  "password": "SecurePassword123!",
  "role": "DONOR",
  "bloodGroup": "O+",
  "district": "Chennai",
  "latitude": 13.0827,
  "longitude": 80.2707
}

// 1. RESPONSE: HTTP 201 Created
{
  "success": true,
  "message": "User registered successfully",
  "data": {
    "user": { "_id": "66fa9b128c1e4a0012bc4401", "name": "Samyak Pon", "role": "DONOR" },
    "token": "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9..."
  }
}

// 2. REQUEST: POST /api/requests (Hospital Emergency Requisition)
// Headers: Authorization: Bearer &lt;Hospital JWT&gt;
{
  "bloodGroup": "O+",
  "component": "WHOLE",
  "unitsNeeded": 2,
  "urgency": "CRITICAL"
}

// 2. RESPONSE: HTTP 201 Created
{
  "success": true,
  "message": "Emergency blood request published and matching donors alerted.",
  "data": {
    "_id": "66fa9c898c1e4a0012bc4555",
    "hospital": "66fa9b128c1e4a0012bc4302",
    "bloodGroup": "O+",
    "component": "WHOLE",
    "unitsNeeded": 2,
    "urgency": "CRITICAL",
    "status": "ACTIVE",
    "potentialDonorsCount": 8
  }
}
    </div>
    """

    # Page 50: Automated Vitest Regression Test Suite
    pages[49] = """
    <h2 class="section-title">SOURCE CODE 9: AUTOMATED INTEGRATION TEST SUITE</h2>
    <p class="justify">
      <strong>File: server/src/tests/donorLocationFlow.test.ts</strong>
    </p>
    <div style="background: #f8fafc; border: 1px solid #cbd5e1; padding: 8px; margin: 8px 0; font-family: 'Courier New', monospace; font-size: 7pt; line-height: 1.18;">
import { describe, it, expect, beforeAll, afterAll } from 'vitest';
import request from 'supertest';
import app from '../app';
import matchingService from '../services/matchingService';

describe('Donor Registration &amp; Geospatial Matching Regression Suite', () => {
  it('TC-01: Should successfully register donor without coordinates using sparse index', async () => {
    const res = await request(app).post('/api/auth/register').send({
      name: 'Test Donor Standby',
      email: `donor_${Date.now()}@example.com`,
      password: 'Password@123',
      role: 'DONOR',
      bloodGroup: 'O+',
      district: 'Chennai'
    });
    expect(res.status).toBe(201);
    expect(res.body.success).toBe(true);
    expect(res.body.data.donorProfile.location).toBeUndefined();
  });

  it('TC-03: Should reject invalid latitude (&gt; 90) with 400 Bad Request', async () => {
    const res = await request(app).post('/api/auth/register').send({
      name: 'Invalid Coords Donor',
      email: `invalid_${Date.now()}@example.com`,
      password: 'Password@123',
      role: 'DONOR',
      bloodGroup: 'A+',
      district: 'Chennai',
      latitude: 95.0,
      longitude: 80.2
    });
    expect(res.status).toBe(400);
    expect(res.body.message).toContain('Latitude must be a valid number between -90 and 90');
  });

  it('TC-06: Should match eligible donor in same district regardless of GPS distance', async () => {
    const matchResult = await matchingService.findMatchingDonors({
      bloodGroup: 'B+',
      component: 'WHOLE',
      hospitalDistrict: 'Chennai',
      hospitalLocation: { type: 'Point', coordinates: [80.2707, 13.0827] }
    } as any);
    expect(matchResult.length).toBeGreaterThan(0);
    expect(matchResult.some(d => d.district === 'Chennai')).toBe(true);
  });
});
    </div>
    """

    # -------------------------------------------------------------------------
    # PAGES 51 to 54: SCREENSHOTS OF THE PROJECT (4 DEDICATED PAGES, 2 FIGURES EACH)
    # The user requested: "leave the screenshots page later I will send you the
    # screenshots and then you attach that later"
    # We construct clean, professional placeholder frames with captions!
    # -------------------------------------------------------------------------

    # Helper function for screenshot placeholder frame
    def render_screenshot_box(fig_id, title, desc):
        return f"""
        <div style="margin-bottom: 25px;">
          <p style="font-weight: bold; font-size: 10.5pt; margin-bottom: 6px; color: #0f172a;">{fig_id}: {title}</p>
          <div style="border: 2px dashed #94a3b8; height: 260px; display: flex; flex-direction: column; align-items: center; justify-content: center; background: #f8fafc; border-radius: 8px;">
            <div style="font-size: 32pt; color: #64748b; margin-bottom: 8px;">🖼️</div>
            <p style="font-weight: bold; font-size: 11pt; color: #334155; margin: 0 0 4px 0;">[ SCREENSHOT PLACEHOLDER — {fig_id} ]</p>
            <p style="font-size: 9pt; color: #64748b; font-style: italic; margin: 0;">Frame reserved for application interface screenshot (User to attach later)</p>
          </div>
          <p style="font-size: 9.5pt; font-style: italic; color: #334155; margin-top: 6px; text-align: center;">
            <strong>{fig_id}:</strong> {desc}
          </p>
        </div>
        """

    # PAGE 51: Screenshots S.1 & S.2
    pages[50] = f"""
    <div style="text-align: center; padding-top: 10px; margin-bottom: 15px;">
      <h1 style="font-size: 13pt; font-weight: bold; text-transform: uppercase;">SCREENSHOTS OF THE PROJECT</h1>
    </div>
    {render_screenshot_box('Figure S.1', 'BloodLink Public Landing &amp; Multi-Role Onboarding Portal', 'Public landing portal with emergency urgency awareness, multi-role registration, and secure credential onboarding.')}
    {render_screenshot_box('Figure S.2', 'Disaster &amp; Operations Command Center — Real-Time Emergency Matching Feed', 'Administrative command center tracking verified hospital statuses, active emergency requests, and system-wide donor mobilization velocity.')}
    """

    # PAGE 52: Screenshots S.3 & S.4
    pages[51] = f"""
    <div style="text-align: center; padding-top: 10px; margin-bottom: 15px;">
      <h1 style="font-size: 13pt; font-weight: bold; text-transform: uppercase;">SCREENSHOTS OF THE PROJECT (CONTINUED)</h1>
    </div>
    {render_screenshot_box('Figure S.3', 'Hospital Emergency Blood Request Creation &amp; Urgency Triage HUD', 'Clinical request authoring console enabling rapid selection of ABO/Rh group, blood component, units needed, and urgency level with live matching preview.')}
    {render_screenshot_box('Figure S.4', 'Voluntary Donor Real-Time Emergency Triage Feed &amp; Route Guidance', 'Mobile-responsive donor triage feed displaying emergency alerts matching donor blood group and district, with single-click acceptance lock.')}
    """

    # PAGE 53: Screenshots S.5 & S.6
    pages[52] = f"""
    <div style="text-align: center; padding-top: 10px; margin-bottom: 15px;">
      <h1 style="font-size: 13pt; font-weight: bold; text-transform: uppercase;">SCREENSHOTS OF THE PROJECT (CONTINUED)</h1>
    </div>
    {render_screenshot_box('Figure S.5', 'Donor Verified Donation History Certificate &amp; Lifetime Impact Metrics', 'Donor profile interface showcasing verified lifetime donation history, readiness toggles, and cryptographic donation confirmation certificates.')}
    {render_screenshot_box('Figure S.6', 'System Administrator Hospital Verification &amp; Regulatory Audit Portal', 'Regulatory administrative console enabling designated health officials to audit hospital operating licenses, verify facility coordinates, and approve medical access.')}
    """

    # PAGE 54: Screenshots S.7 & S.8
    pages[53] = f"""
    <div style="text-align: center; padding-top: 10px; margin-bottom: 15px;">
      <h1 style="font-size: 13pt; font-weight: bold; text-transform: uppercase;">SCREENSHOTS OF THE PROJECT (CONTINUED)</h1>
    </div>
    {render_screenshot_box('Figure S.7', 'Mobile Responsive Emergency Dispatch Interface &amp; Standby HUD', 'Smartphone-optimized donor standby HUD featuring acoustic alert triggers, regional language selector, and immediate acceptance verification.')}
    {render_screenshot_box('Figure S.8', 'System Audit Logs &amp; DPDP Compliance Security Dashboard', 'Immutable audit log viewer documenting sensitive clinical state changes, cryptographic token verifications, and compliance enforcement.')}
    """

    # PAGE 55: REFERENCES (Exact Match to Template)
    pages[54] = raw_pages[54] # Original comprehensive 20 IEEE references

    return pages

def compile_all_formats():
    pages = build_custom_pages()
    total_pages = len(pages)
    print(f"Total pages assembled: {total_pages}")
    assert total_pages == 55, f"Expected 55 pages, got {total_pages}"

    # Generate HTML
    html_parts = [
        "<!DOCTYPE html>",
        "<html lang='en'>",
        "<head>",
        "<meta charset='UTF-8'>",
        "<title>BloodLink: Real-Time Emergency Blood Donor & Hospital Matching Platform - Academic Project Report</title>",
        "<style>",
        "  @page {",
        "    size: A4 portrait;",
        "    margin: 20mm 18mm 20mm 20mm;",
        "  }",
        "  body {",
        "    font-family: 'Times New Roman', Times, serif;",
        "    font-size: 11pt;",
        "    line-height: 1.45;",
        "    color: #000;",
        "    margin: 0;",
        "    padding: 0;",
        "    background-color: #fff;",
        "  }",
        "  .page {",
        "    width: 100%;",
        "    height: 100vh;",
        "    max-height: 100vh;",
        "    box-sizing: border-box;",
        "    page-break-after: always;",
        "    page-break-inside: avoid;",
        "    overflow: hidden;",
        "    position: relative;",
        "    display: flex;",
        "    flex-direction: column;",
        "    justify-content: flex-start;",
        "  }",
        "  .page-num {",
        "    position: absolute;",
        "    bottom: 8px;",
        "    left: 0;",
        "    width: 100%;",
        "    text-align: center;",
        "    font-size: 10pt;",
        "  }",
        "  h1.chapter-title {",
        "    font-size: 14pt;",
        "    font-weight: bold;",
        "    text-align: center;",
        "    margin-top: 15px;",
        "    margin-bottom: 18px;",
        "    text-transform: uppercase;",
        "  }",
        "  h2.section-title {",
        "    font-size: 12pt;",
        "    font-weight: bold;",
        "    margin-top: 14px;",
        "    margin-bottom: 8px;",
        "  }",
        "  h3.subsection-title {",
        "    font-size: 11pt;",
        "    font-weight: bold;",
        "    margin-top: 10px;",
        "    margin-bottom: 6px;",
        "  }",
        "  p.justify {",
        "    text-align: justify;",
        "    text-justify: inter-word;",
        "    margin-top: 0;",
        "    margin-bottom: 10px;",
        "  }",
        "  ul, ol {",
        "    margin-top: 0;",
        "    margin-bottom: 10px;",
        "    padding-left: 24px;",
        "  }",
        "  li {",
        "    margin-bottom: 5px;",
        "    text-align: justify;",
        "  }",
        "  table {",
        "    width: 100%;",
        "    border-collapse: collapse;",
        "    margin-top: 8px;",
        "    margin-bottom: 10px;",
        "  }",
        "  th, td {",
        "    border: 1px solid #475569;",
        "    padding: 4px 6px;",
        "    text-align: left;",
        "    vertical-align: top;",
        "  }",
        "  th {",
        "    background-color: #e2e8f0;",
        "    font-weight: bold;",
        "  }",
        "  .print-btn {",
        "    position: fixed;",
        "    top: 15px;",
        "    right: 15px;",
        "    background-color: #dc2626;",
        "    color: white;",
        "    padding: 10px 18px;",
        "    border-radius: 6px;",
        "    font-weight: bold;",
        "    cursor: pointer;",
        "    box-shadow: 0 4px 6px rgba(0,0,0,0.2);",
        "    z-index: 9999;",
        "  }",
        "  @media print {",
        "    .print-btn { display: none; }",
        "    .page { height: 100vh; max-height: 100vh; page-break-after: always; }",
        "  }",
        "</style>",
        "</head>",
        "<body>",
        "<button class='print-btn' onclick='window.print()'>Print to PDF (55 Pages)</button>"
    ]

    for idx, page_content in enumerate(pages):
        page_num = idx + 1
        page_footer = f"<div class='page-num'>{page_num}</div>"
        html_parts.append(f"<div class='page' id='page-{page_num}'>")
        html_parts.append(page_content)
        html_parts.append(page_footer)
        html_parts.append("</div>\n")

    html_parts.append("</body></html>")
    
    html_path = os.path.abspath('BloodLink_Project_Report.html')
    with open(html_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(html_parts))
    print(f"Successfully generated {html_path}")

    # Generate PDF via Headless Edge/Chrome
    pdf_path = os.path.abspath('BloodLink_Project_Report.pdf')
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"
    ]
    browser_exe = None
    for c in candidates:
        if os.path.exists(c):
            browser_exe = c
            break

    if browser_exe:
        cmd = [
            browser_exe,
            "--headless",
            "--disable-gpu",
            "--run-all-compositor-stages-before-draw",
            f"--print-to-pdf={pdf_path}",
            html_path
        ]
        subprocess.run(cmd, check=True)
        reader = pypdf.PdfReader(pdf_path)
        print(f"PDF Generated successfully! Total pages: {len(reader.pages)}")
    else:
        print("Warning: Browser executable not found for PDF compilation.")

    # Generate DOCX and Markdown
    docx_path = os.path.abspath('BloodLink_Project_Report.docx')
    alt_docx_path = os.path.abspath('BloodLink_Project_Report_Latest.docx')
    md_path = os.path.abspath('BloodLink_Project_Report.md')

    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.7)

    md_content = ["# BLOODLINK: REAL-TIME EMERGENCY BLOOD DONOR & HOSPITAL MATCHING PLATFORM\n"]

    def sanitize_xml(s):
        s = html.unescape(s)
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', s)

    for idx, page_html in enumerate(pages):
        if idx > 0:
            doc.add_page_break()
            md_content.append(f"\n---\n\n<!-- PAGE {idx+1} -->\n")
        
        clean_text = re.sub(r'<table[\s\S]*?</table>', '[TABLE INSERTED]', page_html)
        clean_text = re.sub(r'<[^>]+>', ' ', clean_text)
        clean_text = sanitize_xml(clean_text)
        lines = [l.strip() for l in clean_text.split('\n') if l.strip()]
        
        for raw_line in lines:
            line = sanitize_xml(raw_line)
            if not line:
                continue
            if line.startswith('CHAPTER') or line in ['ABSTRACT', 'ACKNOWLEDGEMENT', 'TABLE OF CONTENTS', 'REFERENCES', 'SCREENSHOTS OF THE PROJECT']:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(14)
                md_content.append(f"\n# {line}\n")
            elif any(line.startswith(x) for x in ['1.', '2.', '3.', '4.', '5.', '6.', '7.', 'Table', 'Figure', 'Use Case', 'Section', 'SOURCE CODE']):
                p = doc.add_paragraph()
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(11)
                md_content.append(f"\n## {line}\n")
            else:
                p = doc.add_paragraph(line)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                md_content.append(f"{line}\n\n")

    try:
        doc.save(docx_path)
        print(f"Successfully saved {docx_path}")
    except PermissionError:
        doc.save(alt_docx_path)
        print(f"File {docx_path} was locked by Word. Successfully saved {alt_docx_path}")

    with open(md_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(md_content))
    print(f"Successfully generated {md_path}")

if __name__ == '__main__':
    compile_all_formats()
